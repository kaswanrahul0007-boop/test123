import csv
from datetime import datetime
from botlib.config import expand
from botlib.tasks import task


def _summarize(csv_path):
    with open(csv_path, newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))
    sums = {}
    for col in (rows[0].keys() if rows else []):
        try:
            sums[col] = sum(float(r[col]) for r in rows if r[col] not in ("", None))
        except ValueError:
            pass
    return {"file": csv_path.name, "rows": len(rows), "sums": sums, "data": rows}


@task("report")
def report(cfg):
    t = cfg["tasks"]["report"]
    src, out = expand(t["source_folder"]), expand(t["output_dir"])
    src.mkdir(parents=True, exist_ok=True)
    out.mkdir(parents=True, exist_ok=True)
    summaries = [_summarize(p) for p in sorted(src.glob("*.csv"))]
    stamp = datetime.now().strftime("%Y-%m-%d_%H%M")
    made = []

    if "xlsx" in t.get("formats", []):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Summary"
        ws.append(["File", "Rows", "Column", "Total"])
        for s in summaries:
            for col, tot in (s["sums"] or {"-": ""}).items():
                ws.append([s["file"], s["rows"], col, tot])
        for s in summaries:
            sh = wb.create_sheet(s["file"][:31])
            if s["data"]:
                sh.append(list(s["data"][0].keys()))
                for r in s["data"]:
                    sh.append(list(r.values()))
        p = out / f"report_{stamp}.xlsx"
        wb.save(p)
        made.append(str(p))

    if "docx" in t.get("formats", []):
        from docx import Document
        d = Document()
        d.add_heading(f"Report {stamp}", 0)
        for s in summaries:
            d.add_heading(s["file"], 2)
            d.add_paragraph(f"{s['rows']} rows")
            for col, tot in s["sums"].items():
                d.add_paragraph(f"{col}: {tot:,.2f}", style="List Bullet")
        p = out / f"report_{stamp}.docx"
        d.save(p)
        made.append(str(p))

    return f"report: {len(summaries)} csv file(s) summarized\n" + "\n".join(made)
